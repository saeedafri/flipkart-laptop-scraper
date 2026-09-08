import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import requests


SCRIPT = Path(__file__).with_name("flipkart_laptop_scraper_exact.py")


class ScraperImportContractTests(unittest.TestCase):
    def test_module_can_be_imported_without_starting_a_scrape(self):
        result = subprocess.run(
            [sys.executable, "-c", "import flipkart_laptop_scraper_exact; print('imported')"],
            cwd=SCRIPT.parent,
            capture_output=True,
            text=True,
            timeout=10,
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(result.stdout.strip(), "imported")

    def test_recent_row_cache_avoids_repeating_completed_product_work(self):
        program = """
import tempfile
from pathlib import Path
from flipkart_scraper.cache import RowCache

with tempfile.TemporaryDirectory() as directory:
    cache = RowCache(Path(directory) / "rows.json", max_age_hours=1, now=lambda: 1000)
    cache.put("product-1", {"Product name": "HP Pavilion"})
    cache.save()
    restored = RowCache(Path(directory) / "rows.json", max_age_hours=1, now=lambda: 1100)
    assert restored.get("product-1") == {"Product name": "HP Pavilion"}
"""
        result = subprocess.run(
            [sys.executable, "-c", program],
            cwd=SCRIPT.parent,
            capture_output=True,
            text=True,
        )
        self.assertEqual(result.returncode, 0, result.stderr)


class ScraperParsingTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import flipkart_laptop_scraper_exact as scraper

        cls.scraper = scraper

    def test_product_page_uses_structured_product_data_for_all_columns(self):
        state = {
            "multiWidgetState": {
                "pageDataResponse": {
                    "seoData": {
                        "schema": [{
                            "@type": "Product",
                            "name": "ASUS TUF Gaming A15 AMD Ryzen 7 - (16 GB/1 TB SSD/Windows 11 Home/4 GB Graphics/NVIDIA GeForce RTX 3050) X Gaming Laptop",
                            "brand": {"name": "ASUS"},
                            "aggregateRating": {"ratingValue": 4.4, "ratingCount": 2479, "reviewCount": 219},
                            "offers": {"price": 70990, "availability": "https://schema.org/InStock"},
                        }]
                    },
                    "pageContext": {
                        "fdpEventTracking": {"events": {"psi": {"ppd": {"finalPrice": 70990, "mrp": 84990}}}}
                    },
                },
                "widgetsData": {"specs": [
                    self.spec("Brand", "ASUS"),
                    self.spec("Series", "TUF Gaming A15"),
                    self.spec("Type", "Gaming Laptop"),
                    self.spec("Suitable For", "Gaming"),
                    self.spec("Processor Brand", "AMD"),
                    self.spec("Processor Name", "Ryzen 7"),
                    self.spec("Processor Variant", "170"),
                    self.spec("RAM", "16 GB"),
                    self.spec("RAM Type", "DDR5"),
                    self.spec("SSD", "Yes"),
                    self.spec("SSD Capacity", "1 TB"),
                    self.spec("Storage Type", "SSD"),
                    self.spec("Dedicated Graphic Memory Capacity", "4 GB"),
                    self.spec("Dedicated Graphic Memory Type", "GDDR6"),
                    self.spec("Graphic Processor", "NVIDIA GeForce RTX 3050"),
                    self.spec("Screen Size", "39.62 cm (15.6 inch)"),
                    self.spec("Touchscreen", "No"),
                    self.spec("Operating System", "Windows 11 Home"),
                ]},
            }
        }
        html = self.page_html(state)

        row = self.scraper.parse_product_page(html, "https://www.flipkart.com/item/p/abc")

        self.assertEqual(list(row), self.scraper.COLUMNS)
        self.assertEqual(row, {
            "Product name": "ASUS TUF Gaming A15",
            "Brand": "ASUS",
            "Laptop type": "Gaming Laptop",
            "Price": 70990,
            "Discount": "16%",
            "MRP": 84990,
            "Rating": 4.4,
            "Review counts": 219,
            "Processor": "AMD Ryzen 7 170",
            "RAM": "16 GB DDR5",
            "Storage": "1024 GB",
            "Storage type": "SSD",
            "Graphics": "Dedicated 4 GB GDDR6",
            "Display size": "39.62 cm (15.6 inch)",
            "GPU": "NVIDIA GeForce RTX 3050",
            "Availability": "In Stock",
            "Touchscreen or not": "No",
            "Operating system": "Windows 11 Home",
            "Laptop use": "Gaming",
        })

    def test_storage_combines_hdd_and_ssd_in_one_consistent_gb_value(self):
        specs = {
            "SSD": "Yes",
            "SSD Capacity": "256 GB",
            "HDD Capacity": "1 TB",
            "Storage Type": "HDD + SSD",
        }
        self.assertEqual(self.scraper.normalise_storage(specs, ""), ("1280 GB", "SSD + HDD"))

    def test_storage_supports_ufs_and_flipkart_emmc_capacity_labels(self):
        self.assertEqual(
            self.scraper.normalise_storage({"Storage Type": "UFS", "UFS Storage Capacity": "256 GB"}, ""),
            ("256 GB", "UFS"),
        )
        self.assertEqual(
            self.scraper.normalise_storage({"Storage Type": "eMMC", "EMMC Storage Capacity": "64 GB"}, ""),
            ("64 GB", "eMMC"),
        )

    def test_processor_does_not_repeat_brand_already_in_processor_name(self):
        specs = {
            "Processor Brand": "MediaTek",
            "Processor Name": "MediaTek Kompanio",
            "Processor Variant": "520",
        }
        self.assertEqual(self.scraper.build_processor(specs, ""), "MediaTek Kompanio 520")

    def test_graphics_distinguishes_integrated_and_dedicated_gpu(self):
        self.assertEqual(self.scraper.build_graphics({}, "", "Intel Graphics Arc"), "Integrated")
        self.assertEqual(self.scraper.build_graphics({}, "", "AMD Radeon Graphics"), "Integrated")
        self.assertEqual(self.scraper.build_graphics({}, "", "NVIDIA GeForce RTX 3050"), "Dedicated")

    def test_spec_extraction_uses_main_product_widget_not_recommendation_noise(self):
        state = {"multiWidgetState": {"widgetsData": {"slots": [
            {"slotData": {"widget": {"data": {"dlsData": {"specs": [
                self.spec("Brand", "ASUS"), self.spec("Series", "Vivobook S14"),
                self.spec("Processor Brand", "Snapdragon"), self.spec("RAM", "16 GB"),
                self.spec("Graphic Processor", "Qualcomm Adreno"), self.spec("Screen Size", "14 inch"),
            ]}}}}},
            {"slotData": {"widget": {"data": {"dlsData": {
                "recommendation": [self.spec("Graphic Processor", "AMD Radeon")]
            }}}}},
        ]}}}
        specs = self.scraper.extract_specs(state)
        self.assertEqual(specs["Graphic Processor"], "Qualcomm Adreno")

    def test_gpu_rejects_wrong_vendor_for_integrated_snapdragon_graphics(self):
        gpu = self.scraper.normalise_gpu("Snapdragon X X1 26 100", "AMD Radeon AMD", "", {})
        self.assertEqual(gpu, "Qualcomm Adreno Integrated Graphics")

    def test_gpu_names_are_canonical_not_repeated_source_fragments(self):
        cases = [
            ("Intel Core i5", "Intel Integrated Intel Integrated UHD", "Intel UHD Graphics"),
            ("AMD Ryzen 5", "AMD Radeon Radeon 610M Graphics", "AMD Radeon 610M"),
            ("Intel Core i7", "NVIDIA GeForce RTX RTX 3050", "NVIDIA GeForce RTX 3050"),
            ("Snapdragon X", "Qualcomm Adreno Qualcomm Adreno Adreno", "Qualcomm Adreno Integrated Graphics"),
        ]
        for processor, raw_gpu, expected in cases:
            with self.subTest(raw_gpu=raw_gpu):
                self.assertEqual(self.scraper.normalise_gpu(processor, raw_gpu, "", {}), expected)

    def test_product_name_removes_bundle_and_marketing_spec_text(self):
        specs = {"Brand": "ASUS", "Series": "Vivobook S14 (2025) with Office 2024 + M365 Basic*, AI PC, Backlit Keyboard"}
        name = self.scraper.clean_product_name("ignored", specs, "Intel Core 5")
        self.assertEqual(name, "ASUS Vivobook S14 (2025)")

    def test_unavailable_and_coming_soon_are_never_marked_in_stock(self):
        for schema_status, visible_text, expected in [
            ("https://schema.org/OutOfStock", "Currently unavailable", "Currently Unavailable"),
            ("https://schema.org/PreOrder", "Coming Soon", "Coming Soon"),
        ]:
            with self.subTest(expected=expected):
                state = {"multiWidgetState": {"pageDataResponse": {"seoData": {"schema": [{
                    "@type": "Product",
                    "name": "HP Pavilion Laptop",
                    "brand": {"name": "HP"},
                    "offers": {"price": 50000, "availability": schema_status},
                }]}}}}
                row = self.scraper.parse_product_page(self.page_html(state, visible_text), "https://example.test/p")
                self.assertEqual(row["Availability"], expected)

    def test_excel_has_real_headers_and_no_dataframe_index_column(self):
        from openpyxl import load_workbook

        row = {column: None for column in self.scraper.COLUMNS}
        row["Product name"] = "HP Pavilion"
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "laptops.xlsx"
            self.scraper.write_excel([row], output)
            sheet = load_workbook(output, read_only=True).active
            headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
            values = [cell.value for cell in next(sheet.iter_rows(min_row=2, max_row=2))]

        self.assertEqual(headers, self.scraper.COLUMNS)
        self.assertEqual(len(values), 19)
        self.assertEqual(values[0], "HP Pavilion")

    def test_excel_removes_rows_that_are_identical_in_all_requested_columns(self):
        row = {column: "value" for column in self.scraper.COLUMNS}
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "laptops.xlsx"
            self.scraper.write_excel([row, row.copy()], output)
            from openpyxl import load_workbook
            sheet = load_workbook(output, read_only=True).active
            self.assertEqual(sheet.max_row, 2)

    def test_403_honours_retry_after_instead_of_retrying_immediately(self):
        class Response:
            def __init__(self, status, retry_after=None):
                self.status_code = status
                self.headers = {"Retry-After": retry_after} if retry_after else {}
                self.text = "x" * 10_001

            def raise_for_status(self):
                if self.status_code >= 400:
                    error = requests.HTTPError(f"{self.status_code} response")
                    error.response = self
                    raise error

        class Session:
            def __init__(self):
                self.responses = iter([Response(403, "12"), Response(200)])

            def get(self, *_args, **_kwargs):
                return next(self.responses)

        waits = []
        with patch("flipkart_scraper.core._session", return_value=Session()), patch(
            "flipkart_scraper.core.time.sleep", side_effect=waits.append
        ):
            result = self.scraper.fetch("https://example.test", retries=2)

        self.assertEqual(result, "x" * 10_001)
        self.assertEqual(waits, [12.0])

    def test_fetch_uses_a_browser_tls_fingerprint(self):
        class Response:
            status_code = 200
            headers = {}
            text = "x" * 10_001

            @staticmethod
            def raise_for_status():
                return None

        class Session:
            def __init__(self):
                self.request_options = None

            def get(self, *_args, **kwargs):
                self.request_options = kwargs
                return Response()

        session = Session()
        with patch("flipkart_scraper.core._session", return_value=session):
            self.scraper.fetch("https://example.test")

        self.assertEqual(session.request_options["impersonate"], "chrome")

    def test_scrape_products_reuses_cached_row_without_network_request(self):
        from flipkart_scraper import core
        from flipkart_scraper.cache import RowCache

        card = {"url": "https://example.test/product-1", "status": "In Stock"}
        row = {column: "cached value" for column in core.COLUMNS}
        with tempfile.TemporaryDirectory() as directory:
            cache = RowCache(Path(directory) / "rows.json", max_age_hours=1)
            cache.put(card["url"], row)
            with patch.object(core, "fetch", side_effect=AssertionError("network should not run")):
                try:
                    rows = core.scrape_products([card], workers=1, cache=cache)
                except TypeError as error:
                    self.fail(f"scrape_products does not support the row cache: {error}")

        self.assertEqual(rows, [row])

    def test_discovery_stops_after_persistent_rate_limit(self):
        from flipkart_scraper import core

        rate_limit_error = getattr(core, "RateLimitError", RuntimeError)
        calls = []

        def blocked(url):
            calls.append(url)
            raise rate_limit_error("Flipkart is still returning HTTP 403")

        with patch.object(core, "fetch", side_effect=blocked):
            with self.assertRaisesRegex(RuntimeError, "HTTP 403"):
                core.discover_products(pages=50, query="laptop", delay=0)

        self.assertEqual(len(calls), 1)

    def test_discovery_reuses_recent_cached_search_page(self):
        from flipkart_scraper import core
        from flipkart_scraper.cache import RowCache

        card = {"id": "ABC", "url": "https://example.test/p", "title": "HP", "status": "In Stock", "text": "HP"}
        with tempfile.TemporaryDirectory() as directory:
            cache = RowCache(Path(directory) / "rows.json", max_age_hours=1)
            cache.put("search:laptop:1", {"cards": [card]})
            with patch.object(core, "fetch", side_effect=AssertionError("network should not run")):
                try:
                    products = core.discover_products(1, "laptop", 0, cache=cache)
                except TypeError as error:
                    self.fail(f"discover_products does not support the search cache: {error}")

        self.assertEqual(products, [card])

    @staticmethod
    def spec(label, value):
        return {"label_0": {"value": {"text": label}}, "label_1": {"value": {"text": [value]}}}

    @staticmethod
    def page_html(state, visible_text=""):
        return (
            "<html><body>" + visible_text + "<script>window.__INITIAL_STATE__ = "
            + json.dumps(state) + ";</script></body></html>"
        )


if __name__ == "__main__":
    unittest.main()
