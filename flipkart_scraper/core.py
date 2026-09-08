"""Scrape Flipkart laptops safely and export the required Excel workbook."""

from __future__ import annotations

import argparse
import json
import re
import threading
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
from typing import Any
from urllib.parse import parse_qs, urlencode, urljoin, urlparse, urlunparse

import pandas as pd
import requests
from bs4 import BeautifulSoup
from openpyxl.styles import Alignment, Font, PatternFill

from .cache import RowCache


BASE_URL = "https://www.flipkart.com"
SEARCH_URL = f"{BASE_URL}/search"
NOT_AVAILABLE = "Not Available"
COLUMNS = [
    "Product name", "Brand", "Laptop type", "Price", "Discount", "MRP",
    "Rating", "Review counts", "Processor", "RAM", "Storage", "Storage type",
    "Graphics", "Display size", "GPU", "Availability", "Touchscreen or not",
    "Operating system", "Laptop use",
]
HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"
    ),
    "Accept-Language": "en-US,en;q=0.9",
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,image/webp,*/*;q=0.8",
    "Referer": "https://www.google.com/",
}
_thread_local = threading.local()


class RateLimitError(RuntimeError):
    """Flipkart kept returning 403/429 after the safe cooldown retries."""


def clean_text(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def first_available(*values: Any) -> Any:
    for value in values:
        cleaned = clean_text(value)
        if value is not None and cleaned and cleaned.lower() not in {
            "na", "n/a", "none", "null", "not applicable", NOT_AVAILABLE.lower(),
        }:
            return value
    return NOT_AVAILABLE


def number(value: Any) -> int | float | None:
    if isinstance(value, (int, float)) and not pd.isna(value):
        return value
    match = re.search(r"\d[\d,]*(?:\.\d+)?", clean_text(value))
    if not match:
        return None
    parsed = float(match.group(0).replace(",", ""))
    return int(parsed) if parsed.is_integer() else parsed


def _text_value(node: Any) -> str:
    if not isinstance(node, dict):
        return ""
    value = node.get("value", node)
    if not isinstance(value, dict) or "text" not in value:
        return ""
    text = value["text"]
    if isinstance(text, list):
        return clean_text(" ".join(str(part) for part in text))
    return clean_text(text)


def extract_specs(state: dict[str, Any]) -> dict[str, str]:
    recognised = {
        "Brand", "Series", "Type", "Suitable For", "Processor Brand", "Processor Name",
        "Processor Variant", "RAM", "RAM Type", "SSD", "SSD Capacity", "HDD Capacity",
        "eMMC Capacity", "EMMC Storage Capacity", "UFS Storage Capacity", "Storage Type",
        "Dedicated Graphic Memory Capacity",
        "Dedicated Graphic Memory Type", "Graphic Processor", "Screen Size", "Touchscreen",
        "Operating System", "Model Number", "Model Name",
    }

    def collect(node: Any) -> dict[str, str]:
        found: dict[str, str] = {}

        def walk(child: Any) -> None:
            if isinstance(child, dict):
                label = _text_value(child.get("label_0"))
                value = _text_value(child.get("label_1")) or _text_value(child.get("label_2"))
                if label and value:
                    found[label.strip()] = value
                for descendant in child.values():
                    walk(descendant)
            elif isinstance(child, list):
                for descendant in child:
                    walk(descendant)

        walk(node)
        return found

    widgets = state.get("multiWidgetState", {}).get("widgetsData", {})
    slots = widgets.get("slots", []) if isinstance(widgets, dict) else []
    candidates = [collect(slot) for slot in slots] if slots else [collect(widgets)]
    return max(candidates, key=lambda candidate: sum(key in recognised for key in candidate), default={})


def extract_initial_state(html: str) -> dict[str, Any]:
    soup = BeautifulSoup(html, "html.parser")
    for script in soup.find_all("script"):
        source = script.string or script.get_text()
        if source and source.lstrip().startswith("window.__INITIAL_STATE__"):
            return json.loads(source.split("=", 1)[1].strip().rstrip(";"))
    return {}


def find_product_schema(state: dict[str, Any]) -> dict[str, Any]:
    schemas = (
        state.get("multiWidgetState", {}).get("pageDataResponse", {})
        .get("seoData", {}).get("schema", [])
    )
    for schema in schemas if isinstance(schemas, list) else [schemas]:
        if isinstance(schema, dict) and schema.get("@type") == "Product":
            return schema
    return {}


def get_price_data(state: dict[str, Any]) -> dict[str, Any]:
    return (
        state.get("multiWidgetState", {}).get("pageDataResponse", {})
        .get("pageContext", {}).get("fdpEventTracking", {}).get("events", {})
        .get("psi", {}).get("ppd", {})
    )


def clean_product_name(title: str, specs: dict[str, str], processor: str) -> str:
    brand = specs.get("Brand", "")
    series = specs.get("Series", "")
    if series:
        series = re.sub(r"\s*\((?:i[3579]|Intel|AMD|Ryzen|Core)[^)]*\)", "", series, flags=re.IGNORECASE)
        series = re.split(r"\s+(?:with|for\s+Creator)\b", series, maxsplit=1, flags=re.IGNORECASE)[0]
        series = re.split(r",\s*(?:AI\s+PC|Copilot|MSO|Office|Backlit|Metal\s+Body)\b", series, maxsplit=1, flags=re.IGNORECASE)[0]
        series = series.rstrip(" ,")
        return series if series.lower().startswith(brand.lower()) else clean_text(f"{brand} {series}")
    name = re.sub(r"^Add to Compare\s*", "", clean_text(title), flags=re.IGNORECASE)
    name = re.split(r"\s+-\s*\(", name, maxsplit=1)[0]
    if processor and processor != NOT_AVAILABLE:
        name = re.sub(rf"\s+{re.escape(processor)}$", "", name, flags=re.IGNORECASE)
    return name or NOT_AVAILABLE


def build_processor(specs: dict[str, str], title: str) -> str:
    brand = clean_text(specs.get("Processor Brand", ""))
    name = clean_text(specs.get("Processor Name", ""))
    variant = clean_text(specs.get("Processor Variant", ""))
    if name:
        processor = name if name.lower().startswith(brand.lower()) else clean_text(f"{brand} {name}")
        if variant and variant.lower() not in processor.lower().split():
            processor = clean_text(f"{processor} {variant}")
        return processor
    if brand:
        return clean_text(f"{brand} {variant}")
    patterns = [
        r"Intel\s+Core\s+(?:Ultra\s+)?(?:i?[3579]|X[579])(?:\s+\d{3,5}[A-Z]{0,3})?",
        r"AMD\s+Ryzen\s+(?:AI\s+)?[3579](?:\s+\d{3,5}[A-Z]{0,3})?",
        r"Apple\s+M[1-5](?:\s+(?:Pro|Max|Ultra))?",
        r"MediaTek\s+Kompanio\s+\d+",
        r"Intel\s+(?:Celeron|Pentium)\s+\w+",
        r"Qualcomm\s+Snapdragon\s+[^\-/()]+",
    ]
    for pattern in patterns:
        match = re.search(pattern, title, re.IGNORECASE)
        if match:
            return clean_text(match.group(0))
    return NOT_AVAILABLE


def build_ram(specs: dict[str, str], title: str) -> str:
    capacity = specs.get("RAM", "")
    ram_type = specs.get("RAM Type", "")
    if capacity:
        return clean_text(f"{capacity} {ram_type}")
    match = re.search(r"(\d+\s*GB)\s*(LPDDR\w*|DDR\w*)?\s*RAM", title, re.IGNORECASE)
    return clean_text(" ".join(part for part in match.groups() if part)) if match else NOT_AVAILABLE


def _capacity_in_gb(value: str) -> int | float | None:
    match = re.search(r"(\d+(?:\.\d+)?)\s*(TB|GB)", clean_text(value), re.IGNORECASE)
    if not match:
        return None
    amount = float(match.group(1))
    gb = amount * 1024 if match.group(2).upper() == "TB" else amount
    return int(gb) if gb.is_integer() else round(gb, 2)


def normalise_storage(specs: dict[str, str], title: str) -> tuple[str, str]:
    capacities: list[int | float] = []
    storage_types: list[str] = []
    capacity_fields = (
        ("SSD Capacity", "SSD"), ("HDD Capacity", "HDD"),
        ("eMMC Capacity", "eMMC"), ("EMMC Storage Capacity", "eMMC"),
        ("UFS Storage Capacity", "UFS"),
    )
    for key, storage_type in capacity_fields:
        capacity = _capacity_in_gb(specs.get(key, ""))
        if capacity is not None:
            capacities.append(capacity)
            storage_types.append(storage_type)
    if not capacities:
        pattern = r"(\d+(?:\.\d+)?)\s*(TB|GB)\s*(SSD|HDD|eMMC)(?:\s+Storage)?"
        for match in re.finditer(pattern, title, re.IGNORECASE):
            capacity = _capacity_in_gb(f"{match.group(1)} {match.group(2)}")
            if capacity is not None:
                capacities.append(capacity)
                storage_types.append("eMMC" if match.group(3).lower() == "emmc" else match.group(3).upper())
    declared_type = clean_text(specs.get("Storage Type", ""))
    for storage_type in ("SSD", "HDD", "eMMC", "UFS"):
        if re.search(rf"\b{re.escape(storage_type)}\b", declared_type, re.IGNORECASE):
            if storage_type not in storage_types:
                storage_types.append(storage_type)
    if not capacities:
        return NOT_AVAILABLE, " + ".join(storage_types) or NOT_AVAILABLE
    total = sum(capacities)
    amount = int(total) if float(total).is_integer() else round(total, 2)
    ordered_types = [kind for kind in ("SSD", "HDD", "eMMC", "UFS") if kind in storage_types]
    return f"{amount} GB", " + ".join(ordered_types) or NOT_AVAILABLE


def build_graphics(specs: dict[str, str], title: str, gpu: str) -> str:
    capacity = specs.get("Dedicated Graphic Memory Capacity", "")
    memory_type = specs.get("Dedicated Graphic Memory Type", "")
    if not capacity:
        match = re.search(r"(\d+\s*GB)\s+Graphics", title, re.IGNORECASE)
        capacity = match.group(1) if match else ""
    is_discrete_gpu = bool(re.search(r"\b(?:NVIDIA|GeForce|RTX|GTX|Radeon\s+RX)\b", gpu, re.IGNORECASE))
    if capacity or is_discrete_gpu:
        return clean_text(f"Dedicated {capacity} {memory_type}")
    if gpu != NOT_AVAILABLE:
        integrated = bool(re.search(
            r"integrated|\bintel\b(?!.*\b(?:nvidia|geforce)\b)|radeon\s+graphics$|\bmediatek\b",
            gpu,
            re.IGNORECASE,
        ))
        return "Integrated" if integrated else "Dedicated"
    return NOT_AVAILABLE


def normalise_gpu(processor: str, raw_gpu: str, title: str, specs: dict[str, str]) -> str:
    gpu = clean_text(first_available(raw_gpu, _gpu_from_title(title)))
    gpu = gpu.replace("™", "").replace("®", "")
    has_dedicated = bool(
        specs.get("Dedicated Graphic Memory Capacity")
        or re.search(r"\b(?:NVIDIA|GeForce|RTX|GTX|Radeon\s+RX)\b", title, re.IGNORECASE)
    )
    platform = processor.lower()
    if "snapdragon" in platform or "qualcomm" in platform:
        return "Qualcomm Adreno Integrated Graphics"
    if not has_dedicated and "mediatek" in platform and "mali" not in gpu.lower():
        return "ARM Mali Integrated Graphics"
    nvidia = re.search(
        r"(?:NVIDIA\s+)?GeForce\s+(RTX|GTX|MX)\s*(?:(?:RTX|GTX|MX)\s*)?(\d{3,4})(\s*Ti)?(?:\s+with\s+Max-Q\s+Design)?",
        gpu,
        re.IGNORECASE,
    )
    if nvidia:
        suffix = " Ti" if nvidia.group(3) else ""
        return f"NVIDIA GeForce {nvidia.group(1).upper()} {nvidia.group(2)}{suffix}"
    amd_rx = re.search(r"Radeon(?:\s+Radeon)?\s+RX\s*(\d{3,4}[A-Z]{0,3})", gpu, re.IGNORECASE)
    if amd_rx:
        return f"AMD Radeon RX {amd_rx.group(1).upper()}"
    if "radeon" in gpu.lower() or "amd" in platform or "ryzen" in platform:
        amd_model = re.search(
            r"Radeon(?:\s+(?:Integrated\s+AMD\s+Radeon|Radeon))?\s+(?!AMD\b|Graphics\b|Ryzen\b)(R?\d{1,4}[A-Z]{0,3})\b",
            gpu,
            re.IGNORECASE,
        )
        return f"AMD Radeon {amd_model.group(1).upper()}" if amd_model else "AMD Radeon Integrated Graphics"
    if "intel" in gpu.lower() or "intel" in platform:
        arc = re.search(r"\bArc(?:\s+Graphics)?\s*([AB]?\d{3})?\b", gpu, re.IGNORECASE)
        if arc:
            model = f" {arc.group(1).upper()}" if arc.group(1) else ""
            return f"Intel Arc{model} Graphics"
        if re.search(r"Iris\s*Xe", gpu, re.IGNORECASE):
            return "Intel Iris Xe Graphics"
        uhd = re.search(r"\bUHD(?:\s+Graphics)?\s*(\d{3})?\b", gpu, re.IGNORECASE)
        if uhd:
            model = f" {uhd.group(1)}" if uhd.group(1) else ""
            return f"Intel UHD{model} Graphics"
        intel_hd = re.search(r"\bHD(?:\s+Graphics)?\s*(\d{3,4})?\b", gpu, re.IGNORECASE)
        if intel_hd:
            model = f" {intel_hd.group(1)}" if intel_hd.group(1) else ""
            return f"Intel HD{model} Graphics"
        return "Intel Integrated Graphics"
    if "mali" in gpu.lower():
        mali = re.search(r"Mali[- ]?[A-Z0-9]+(?:\s+\w+)?", gpu, re.IGNORECASE)
        return f"ARM {clean_text(mali.group(0))} Integrated Graphics" if mali else "ARM Mali Integrated Graphics"
    if gpu.lower() in {"", "na", "n/a", NOT_AVAILABLE.lower()}:
        if "apple" in platform or re.search(r"\bM[1-5]\b", processor):
            return "Apple Integrated GPU"
        if "intel" in platform:
            return "Intel Integrated Graphics"
        if "amd" in platform or "ryzen" in platform:
            return "AMD Radeon Integrated Graphics"
        return NOT_AVAILABLE
    return gpu


def availability(schema_status: str, page_text: str, listing_status: str = "") -> str:
    combined = clean_text(f"{page_text} {listing_status}")
    if re.search(r"\bcoming\s+soon\b", combined, re.IGNORECASE):
        return "Coming Soon"
    if re.search(r"\bcurrently\s+unavailable\b", combined, re.IGNORECASE):
        return "Currently Unavailable"
    if re.search(r"\b(?:sold\s+out|out\s+of\s+stock)\b", combined, re.IGNORECASE):
        return "Out of Stock"
    status = clean_text(schema_status).rsplit("/", 1)[-1].lower()
    return {
        "instock": "In Stock", "outofstock": "Currently Unavailable",
        "preorder": "Coming Soon", "presale": "Coming Soon",
        "backorder": "Back Order", "discontinued": "Discontinued",
    }.get(status, NOT_AVAILABLE)


def guess_laptop_type(title: str) -> str:
    lowered = title.lower()
    if any(word in lowered for word in ("gaming", "rog", "predator", "legion", "victus", "tuf")):
        return "Gaming Laptop"
    if any(word in lowered for word in ("2 in 1", "2-in-1", "convertible")):
        return "2 in 1 Laptop"
    if any(word in lowered for word in ("thin and light", "ultrabook")):
        return "Thin and Light Laptop"
    if "chromebook" in lowered:
        return "Chromebook"
    return "Laptop"


def parse_product_page(html: str, product_url: str = "", listing_status: str = "") -> dict[str, Any]:
    state = extract_initial_state(html)
    schema = find_product_schema(state)
    specs = extract_specs(state)
    price_data = get_price_data(state)
    soup = BeautifulSoup(html, "html.parser")
    page_text = soup.get_text(" ", strip=True)
    title = clean_text(schema.get("name"))
    if not title and soup.title:
        title = re.split(r"\s+Rs\.|\s+Price in India", soup.title.get_text(" ", strip=True), maxsplit=1)[0]
    brand_data = schema.get("brand", {})
    schema_brand = brand_data.get("name", "") if isinstance(brand_data, dict) else brand_data
    if schema_brand and "Brand" not in specs:
        specs["Brand"] = clean_text(schema_brand)
    processor = build_processor(specs, title)
    storage, storage_type = normalise_storage(specs, title)
    raw_gpu = first_available(specs.get("Graphic Processor"), _gpu_from_title(title))
    gpu = normalise_gpu(processor, raw_gpu, title, specs)
    offers = schema.get("offers", {}) if isinstance(schema.get("offers"), dict) else {}
    price = number(first_available(price_data.get("finalPrice"), price_data.get("fsp"), offers.get("price")))
    mrp = number(price_data.get("mrp"))
    discount = NOT_AVAILABLE
    if price is not None and mrp is not None and mrp >= price and mrp > 0:
        discount = f"{round((mrp - price) * 100 / mrp)}%"
    rating_data = schema.get("aggregateRating", {}) if isinstance(schema.get("aggregateRating"), dict) else {}
    row = {
        "Product name": clean_product_name(title, specs, processor),
        "Brand": first_available(specs.get("Brand"), schema_brand, title.split(maxsplit=1)[0] if title else ""),
        "Laptop type": first_available(specs.get("Type"), guess_laptop_type(title)),
        "Price": price if price is not None else NOT_AVAILABLE,
        "Discount": discount,
        "MRP": mrp if mrp is not None else NOT_AVAILABLE,
        "Rating": first_available(number(rating_data.get("ratingValue")), _rating_from_text(page_text)),
        "Review counts": first_available(number(rating_data.get("reviewCount")), _reviews_from_text(page_text)),
        "Processor": processor,
        "RAM": build_ram(specs, title),
        "Storage": storage,
        "Storage type": storage_type,
        "Graphics": build_graphics(specs, title, gpu),
        "Display size": first_available(specs.get("Screen Size"), _display_from_text(page_text)),
        "GPU": gpu,
        "Availability": availability(offers.get("availability", ""), page_text, listing_status),
        "Touchscreen or not": first_available(specs.get("Touchscreen"), _touchscreen_from_title(title)),
        "Operating system": first_available(specs.get("Operating System"), _os_from_title(title)),
        "Laptop use": first_available(specs.get("Suitable For"), _laptop_use(title)),
    }
    return {column: row[column] for column in COLUMNS}


def _gpu_from_title(title: str) -> str:
    patterns = [
        r"NVIDIA\s+GeForce\s+(?:RTX|GTX|MX)\s*\w+(?:\s*Ti)?",
        r"AMD\s+Radeon\s+(?:RX\s*)?\w+(?:\s*Graphics)?",
        r"Intel\s+(?:Arc|Iris Xe|UHD)\s*Graphics?",
        r"Apple\s+\d+-core\s+GPU",
    ]
    for pattern in patterns:
        match = re.search(pattern, title, re.IGNORECASE)
        if match:
            return clean_text(match.group(0))
    return NOT_AVAILABLE


def _rating_from_text(text: str) -> float | str:
    match = re.search(r"\b([0-5](?:\.\d)?)\s*(?:\||★)", text)
    return float(match.group(1)) if match else NOT_AVAILABLE


def _reviews_from_text(text: str) -> int | str:
    match = re.search(r"([\d,]+)\s+Reviews?", text, re.IGNORECASE)
    return int(match.group(1).replace(",", "")) if match else NOT_AVAILABLE


def _display_from_text(text: str) -> str:
    match = re.search(r"\d+(?:\.\d+)?\s*cm\s*\(\s*\d+(?:\.\d+)?\s*[Ii]nch\s*\)", text)
    return clean_text(match.group(0)) if match else NOT_AVAILABLE


def _touchscreen_from_title(title: str) -> str:
    return "Yes" if re.search(r"touch|2[ -]in[ -]1|convertible", title, re.IGNORECASE) else NOT_AVAILABLE


def _os_from_title(title: str) -> str:
    match = re.search(r"Windows\s+1[01](?:\s+(?:Home|Pro))?|Chrome\s*OS|macOS|DOS|Ubuntu", title, re.IGNORECASE)
    return clean_text(match.group(0)) if match else NOT_AVAILABLE


def _laptop_use(title: str) -> str:
    return {
        "Gaming Laptop": "Gaming", "Chromebook": "Everyday Use",
        "Thin and Light Laptop": "Travel & Business", "2 in 1 Laptop": "Everyday Use",
    }.get(guess_laptop_type(title), "Everyday Use")


def canonical_product_url(href: str) -> str:
    absolute = urljoin(BASE_URL, href)
    parsed = urlparse(absolute)
    pid = parse_qs(parsed.query).get("pid", [""])[0]
    query = urlencode({"pid": pid}) if pid else ""
    return urlunparse((parsed.scheme, parsed.netloc, parsed.path, "", query, ""))


def parse_search_page(html: str) -> list[dict[str, str]]:
    soup = BeautifulSoup(html, "html.parser")
    products: list[dict[str, str]] = []
    for card in soup.select("div[data-id]"):
        link = card.select_one("a[href*='/p/']")
        if not link:
            continue
        title_node = card.select_one(".RG5Slk")
        image = card.select_one("img[alt]")
        title = clean_text(title_node.get_text(" ", strip=True) if title_node else image.get("alt", "") if image else "")
        card_text = card.get_text(" ", strip=True)
        products.append({
            "id": clean_text(card.get("data-id")),
            "url": canonical_product_url(link.get("href", "")),
            "title": title,
            "status": availability("", card_text),
            "text": card_text,
        })
    return products


def _session() -> requests.Session:
    if not hasattr(_thread_local, "session"):
        session = requests.Session()
        session.headers.update(HEADERS)
        _thread_local.session = session
    return _thread_local.session


def fetch(url: str, retries: int = 3, timeout: int = 30) -> str:
    last_error: Exception | None = None
    last_status: int | None = None
    for attempt in range(retries):
        try:
            response = _session().get(url, timeout=timeout)
            response.raise_for_status()
            if len(response.text) < 10_000:
                raise RuntimeError(f"unexpected short response ({len(response.text)} bytes)")
            return response.text
        except (requests.RequestException, RuntimeError) as error:
            last_error = error
            if attempt + 1 < retries:
                response = getattr(error, "response", None)
                status = getattr(response, "status_code", None)
                last_status = status
                if status in {403, 429}:
                    retry_after = response.headers.get("Retry-After")
                    wait = float(retry_after) if retry_after and retry_after.isdigit() else 15.0 * (2**attempt)
                    print(f"Flipkart rate limit ({status}); cooling down for {wait:.0f}s")
                else:
                    wait = 1.5 * (attempt + 1)
                time.sleep(wait)
    if last_status in {403, 429}:
        raise RateLimitError(
            f"Flipkart is still returning HTTP {last_status} after cooldown retries. "
            "The run was stopped to avoid a 50-page failure loop. Wait, then rerun; "
            "recent completed product rows will be reused from cache."
        )
    raise RuntimeError(f"Could not fetch {url}: {last_error}")


def discover_products(
    pages: int,
    query: str,
    delay: float,
    cache: RowCache | None = None,
) -> list[dict[str, str]]:
    discovered: dict[str, dict[str, str]] = {}
    for page in range(1, pages + 1):
        cache_key = f"search:{query}:{page}"
        cached = cache.get(cache_key) if cache else None
        cards = cached.get("cards") if cached and isinstance(cached.get("cards"), list) else None
        if cards is None:
            url = f"{SEARCH_URL}?{urlencode({'q': query, 'page': page})}"
            try:
                cards = parse_search_page(fetch(url))
            except RateLimitError:
                raise
            except RuntimeError as error:
                print(f"Search page {page}/{pages} failed: {error}")
                continue
            if cache:
                cache.put(cache_key, {"cards": cards})
                cache.save()
            if delay:
                time.sleep(delay)
        else:
            print(f"Search page {page}/{pages}: reused recent cache")
        if not cards:
            print(f"Search page {page}/{pages}: no products; stopping")
            break
        for card in cards:
            discovered.setdefault(card["id"] or card["url"], card)
        print(f"Search page {page}/{pages}: {len(cards)} cards, {len(discovered)} unique products")
    return list(discovered.values())


def _fallback_row(card: dict[str, str]) -> dict[str, Any]:
    fake_state = {"multiWidgetState": {"pageDataResponse": {"seoData": {"schema": [{
        "@type": "Product", "name": card.get("title", ""),
    }]}}}}
    html = (
        f"<html><body>{card.get('text', '')}"
        f"<script>window.__INITIAL_STATE__ = {json.dumps(fake_state)};</script></body></html>"
    )
    return parse_product_page(html, card.get("url", ""), card.get("status", ""))


def scrape_products(
    cards: list[dict[str, str]],
    workers: int,
    cache: RowCache | None = None,
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any] | None] = [None] * len(cards)
    pending: list[tuple[int, dict[str, str]]] = []

    # Reuse only complete, recent rows. Price and availability therefore never
    # stay cached longer than --cache-hours.
    for index, card in enumerate(cards):
        cached = cache.get(card["url"]) if cache else None
        if cached and list(cached) == COLUMNS:
            rows[index] = cached
        else:
            pending.append((index, card))
    cached_count = len(cards) - len(pending)
    if cached_count:
        print(f"Product details: reused {cached_count} recent cached rows")

    def scrape_one(index: int, card: dict[str, str]) -> tuple[int, dict[str, Any], str | None]:
        try:
            return index, parse_product_page(fetch(card["url"]), card["url"], card.get("status", "")), None
        except Exception as error:
            return index, _fallback_row(card), str(error)

    with ThreadPoolExecutor(max_workers=workers) as pool:
        futures = [pool.submit(scrape_one, index, card) for index, card in pending]
        completed = cached_count
        fallback_count = 0
        for future in as_completed(futures):
            index, row, error = future.result()
            rows[index] = row
            completed += 1
            if error:
                fallback_count += 1
                print(f"Detail fallback for product {index + 1}: {error}")
            elif cache:
                cache.put(cards[index]["url"], row)
                if completed % 10 == 0:
                    cache.save()
            if completed % 10 == 0 or completed == len(cards):
                print(f"Product details: {completed}/{len(cards)} ({fallback_count} fallbacks)")
    if cache:
        cache.save()
    return [row for row in rows if row is not None]


def write_excel(rows: list[dict[str, Any]], output: str | Path) -> Path:
    output_path = Path(output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    frame = pd.DataFrame(rows, columns=COLUMNS)
    frame.drop_duplicates(inplace=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        frame.to_excel(writer, sheet_name="Laptops", index=False, header=True)
        sheet = writer.sheets["Laptops"]
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        header_fill = PatternFill("solid", fgColor="1F4E78")
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        widths = {
            "A": 38, "B": 14, "C": 22, "D": 13, "E": 12, "F": 13, "G": 10,
            "H": 15, "I": 26, "J": 18, "K": 15, "L": 18, "M": 26, "N": 24,
            "O": 30, "P": 23, "Q": 20, "R": 23, "S": 22,
        }
        for column, width in widths.items():
            sheet.column_dimensions[column].width = width
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for column in ("D", "F"):
            for cell in sheet[column][1:]:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = '₹#,##0'
        for cell in sheet["G"][1:]:
            if isinstance(cell.value, (int, float)):
                cell.number_format = "0.0"
    return output_path


def validate_rows(rows: list[dict[str, Any]]) -> None:
    if not rows:
        raise RuntimeError("No products were scraped; workbook was not written")
    valid_statuses = {
        "In Stock", "Currently Unavailable", "Coming Soon", "Out of Stock",
        "Back Order", "Discontinued", NOT_AVAILABLE,
    }
    for index, row in enumerate(rows, start=1):
        if list(row) != COLUMNS:
            raise RuntimeError(f"Row {index} does not match the required 19-column contract")
        if row["Availability"] not in valid_statuses:
            raise RuntimeError(f"Row {index} has invalid availability: {row['Availability']}")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--pages", type=int, default=50, help="search-result pages to scan (default: 50)")
    parser.add_argument("--query", default="laptop", help="Flipkart search query (default: laptop)")
    parser.add_argument("--workers", type=int, default=6, help="parallel product-detail requests (default: 6)")
    parser.add_argument("--delay", type=float, default=0.8, help="delay between search pages in seconds")
    parser.add_argument("--output", default="output/flipkart_laptops.xlsx", help="output .xlsx path")
    parser.add_argument("--cache", default=".flipkart_laptop_cache.json", help="resume-cache path")
    parser.add_argument("--cache-hours", type=float, default=1, help="reuse rows newer than this many hours")
    args = parser.parse_args()
    if args.pages < 1 or args.workers < 1 or args.cache_hours < 0:
        parser.error("--pages and --workers must be positive; --cache-hours cannot be negative")
    return args


def main() -> int:
    args = parse_args()
    cache = RowCache(args.cache, max_age_hours=args.cache_hours)
    print(f"Discovering '{args.query}' across up to {args.pages} search pages")
    cards = discover_products(args.pages, args.query, args.delay, cache=cache)
    rows = scrape_products(cards, args.workers, cache=cache)
    validate_rows(rows)
    output = write_excel(rows, args.output)
    print(f"Wrote {len(rows)} parsed rows to {output.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
